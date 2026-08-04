"""Export App Owner Handover and TDD markdown to Word (.docx)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def _add_runs_with_code(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def markdown_to_docx(md_path: Path, out_path: Path, title: str) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = head.add_run(title)
    hr.bold = True
    hr.font.size = Pt(16)
    hr.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"Source: {md_path.as_posix()}")
    sr.italic = True
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    # skip first H1 (duplicated as title)
    if lines and lines[0].startswith("# "):
        i = 1

    table_buf: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buf
        if not table_buf:
            return
        rows = [r for r in table_buf if not all(set(c.strip()) <= {"-", ":"} for c in r)]
        table_buf = []
        if not rows:
            return
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci in range(cols):
                cell = t.rows[ri].cells[ci]
                cell.text = row[ci] if ci < len(row) else ""
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)
                        if ri == 0:
                            run.bold = True
        doc.add_paragraph()

    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and "|" in line[1:]:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            table_buf.append(cells)
            i += 1
            continue
        flush_table()

        if not line.strip():
            i += 1
            continue
        if line.startswith("---"):
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(line[3:].strip())
            r.bold = True
            r.font.size = Pt(13)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(line[4:].strip())
            r.bold = True
            r.font.size = Pt(11)
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines) and lines[i].startswith("```"):
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code_lines))
            r.font.name = "Consolas"
            r.font.size = Pt(8)
            continue
        if line.lstrip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_code(p, line.lstrip()[2:])
            i += 1
            continue
        if re.match(r"^\d+\.\s", line.lstrip()):
            p = doc.add_paragraph(style="List Number")
            text = re.sub(r"^\d+\.\s+", "", line.lstrip())
            _add_runs_with_code(p, text)
            i += 1
            continue

        p = doc.add_paragraph()
        _add_runs_with_code(p, line)
        i += 1

    flush_table()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path.resolve()


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    h = markdown_to_docx(
        root / "docs/governance/app_owner_handover.md",
        root / "docs/evidence/app_owner_handover.docx",
        "App Owner Handover — CIFAR-10 CNN Platform",
    )
    t = markdown_to_docx(
        root / "docs/architecture/technical_design_document.md",
        root / "docs/evidence/technical_design_document.docx",
        "Technical Design Document — CIFAR-10 CNN Platform",
    )
    print(h)
    print(t)


if __name__ == "__main__":
    main()
